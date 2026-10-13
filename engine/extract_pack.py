#!/usr/bin/env python3
"""Extract the v3 Excel pack into engine/data JSON files."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DEFAULT_XLSX = (
    ROOT.parent
    / "docs"
    / "scoring-pack"
    / "DreamStoneHR_HR_Health_Check_Risk_and_Action_Build_Pack_v3.xlsx"
)
OUT = ROOT / "data"


def sheet_rows(wb, name: str, header_row: int = 5) -> list[dict]:
    ws = wb[name]
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    while headers and headers[-1] is None:
        headers.pop()
    rows: list[dict] = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if all(v is None or v == "" for v in vals):
            continue
        row = {}
        for h, v in zip(headers, vals):
            if h is None:
                continue
            row[str(h).strip()] = v
        rows.append(row)
    return rows


def extract(xlsx: Path = DEFAULT_XLSX) -> Path:
    wb = load_workbook(xlsx, data_only=True)
    questions = sheet_rows(wb, "Questions")
    answers = sheet_rows(wb, "Answers")
    outcomes = sheet_rows(wb, "Answer Outcomes")
    actions = sheet_rows(wb, "Action Library")
    report_copy = sheet_rows(wb, "Report Copy")
    tests = sheet_rows(wb, "Test Cases")

    ws = wb["Scoring"]
    params = {
        "max_answer_points": ws["B6"].value,
        "lower_band": ws["B7"].value,
        "upper_band": ws["B8"].value,
        "critical_trigger_maximum": ws["B9"].value,
        "uncertainty_threshold": ws["B10"].value,
        "unknown_scoring_treatment": ws["B11"].value,
        "band_labels": ws["B12"].value,
    }
    domains = []
    for r in range(22, 34):
        code = ws.cell(r, 1).value
        if not code or code == "TOTAL":
            continue
        domains.append(
            {
                "code": code,
                "name": ws.cell(r, 2).value,
                "weight_pct": ws.cell(r, 3).value,
                "core_count": ws.cell(r, 4).value,
                "conditional_count": ws.cell(r, 5).value,
                "core_maximum": ws.cell(r, 6).value,
                "full_catalogue_maximum": ws.cell(r, 7).value,
            }
        )

    ws = wb["Score Sandbox"]
    sandbox = []
    for r in range(26, ws.max_row + 1):
        qid = ws.cell(r, 1).value
        if not qid:
            continue
        sandbox.append(
            {
                "question_id": qid,
                "domain": ws.cell(r, 2).value,
                "eligible": ws.cell(r, 6).value,
                "response": ws.cell(r, 7).value,
            }
        )

    pack = {
        "version": "3.0",
        "source": xlsx.name,
        "params": params,
        "domains": domains,
        "risk_ranks": {
            0: "Unverified",
            1: "Low",
            2: "Medium",
            3: "High",
            4: "Critical",
        },
        "questions": questions,
        "answers": answers,
        "outcomes": outcomes,
        "actions": actions,
        "report_copy": report_copy,
        "sandbox_demo": sandbox,
        "test_cases": tests,
    }

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "pack.json").write_text(json.dumps(pack, indent=2, default=str))
    for key in (
        "questions",
        "answers",
        "outcomes",
        "actions",
        "report_copy",
        "domains",
        "sandbox_demo",
    ):
        (OUT / f"{key}.json").write_text(json.dumps(pack[key], indent=2, default=str))
    (OUT / "params.json").write_text(json.dumps(params, indent=2, default=str))
    return OUT / "pack.json"


if __name__ == "__main__":
    path = extract()
    print(f"Wrote {path}")
